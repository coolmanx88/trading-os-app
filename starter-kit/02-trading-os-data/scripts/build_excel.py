from __future__ import annotations
import json
from pathlib import Path
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference

ROOT = Path(__file__).resolve().parents[1]
STATE_PATH = ROOT / "data" / "state.json"
OUT_PATH = ROOT / "exports" / "Trading_Journal.xlsx"

INSTRUMENTS = {
    "MNQ": {"point_value": 2.0, "fee_per_contract_side": 0.95},
    "NQ": {"point_value": 20.0, "fee_per_contract_side": 0.0},
}


def replay_trade(t):
    qty = 0.0
    avg = 0.0
    gross = 0.0
    costs = 0.0
    net = 0.0
    timeline = []
    cfg = INSTRUMENTS.get(t.get("instrument"), {"point_value": 0, "fee_per_contract_side": 0})
    pv = cfg["point_value"]
    fee_per_side = max(0.0, float(t.get("feePerContractSide", cfg.get("fee_per_contract_side", 0)) or 0))
    events = sorted(t.get("events", []), key=lambda x: x.get("timestamp", ""))
    for ev in events:
        q = float(ev.get("qtyPerAccount", 0) or 0)
        p = float(ev.get("price", 0) or 0)
        eg = ec = en = 0.0
        if ev.get("type") in ("INITIAL", "ADD"):
            new_qty = qty + q
            avg = ((avg * qty) + (p * q)) / new_qty if new_qty else 0.0
            qty = new_qty
        elif ev.get("type") in ("REDUCE", "CLOSE"):
            close_qty = min(q, qty)
            pts = (p - avg) if t.get("direction") == "LONG" else (avg - p)
            eg = pts * close_qty * pv
            ec = fee_per_side * close_qty * 2
            en = eg - ec
            gross += eg
            costs += ec
            net += en
            qty -= close_qty
            if qty <= 0:
                qty = 0.0
                avg = 0.0
        timeline.append({**ev, "qtyAfter": qty, "avgEntryAfter": avg, "eventGross": eg, "eventCosts": ec, "eventRealized": en, "grossAfter": gross, "costsAfter": costs, "realizedAfter": net})
    account_count = sum(len(a.get("accountIds", [])) for a in t.get("allocations", []))
    return {
        "qty": qty,
        "avg": avg,
        "gross": gross,
        "costs": costs,
        "realized": net,
        "account_count": account_count,
        "portfolio_gross": gross * account_count,
        "portfolio_costs": costs * account_count,
        "portfolio_pnl": net * account_count,
        "timeline": timeline,
    }


def account_opening_balance(a):
    try:
        explicit = float(a.get("openingBalance"))
        if explicit > 0:
            return explicit
    except (TypeError, ValueError):
        pass
    return float(a.get("sizeK", 0) or 0) * 1000.0


def account_adjustments(a):
    total = 0.0
    for item in a.get("balanceAdjustments", []) or []:
        try:
            total += float(item.get("amount", 0) or 0)
        except (TypeError, ValueError):
            pass
    return total


def account_trade_pnl(account_id, trades, replay_cache):
    total = 0.0
    for t in trades:
        if any(account_id in (alloc.get("accountIds", []) or []) for alloc in t.get("allocations", []) or []):
            total += replay_cache[t.get("id")]["realized"]
    return total


def style_sheet(ws, freeze="A2"):
    header_fill = PatternFill("solid", fgColor="111827")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="334155")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=thin)
    if freeze:
        ws.freeze_panes = freeze
    ws.sheet_view.rightToLeft = True
    for col in ws.columns:
        width = 10
        for cell in col:
            if cell.value is not None:
                width = max(width, min(34, len(str(cell.value)) + 2))
        ws.column_dimensions[get_column_letter(col[0].column)].width = width


def add_rows(ws, headers, rows):
    ws.append(headers)
    for row in rows:
        ws.append(row)
    style_sheet(ws)


def build():
    state = json.loads(STATE_PATH.read_text(encoding="utf-8"))
    wb = Workbook()
    wb.remove(wb.active)

    company_map = {c["id"]: c for c in state.get("companies", [])}
    trades = state.get("trades", [])
    replay_cache = {t.get("id"): replay_trade(t) for t in trades}

    trade_rows = []
    event_rows = []
    allocation_rows = []
    daily = defaultdict(lambda: {"trades": 0, "closed": 0, "portfolio_pnl": 0.0})
    instrument_stats = defaultdict(lambda: {"trades": 0, "pnl": 0.0})

    for t in trades:
        r = replay_cache[t.get("id")]
        trade_rows.append([
            t.get("id"), t.get("date"), t.get("instrument"), t.get("direction"), t.get("status"),
            t.get("targetPerAccount"), t.get("bufferPoints"), t.get("stopPrice"), r["account_count"],
            r["gross"], r["costs"], r["realized"], r["portfolio_gross"], r["portfolio_costs"], r["portfolio_pnl"], t.get("createdAt"), t.get("closedAt"),
            t.get("closeReason"), t.get("manualReason"),
            (t.get("review") or {}).get("adherence"), (t.get("review") or {}).get("mainError"),
            (t.get("review") or {}).get("lesson"), (t.get("review") or {}).get("reviewedAt"), t.get("notes")
        ])
        d = daily[t.get("date")]
        d["trades"] += 1
        if t.get("status") == "Closed":
            d["closed"] += 1
            d["portfolio_pnl"] += r["portfolio_pnl"]
            instrument_stats[t.get("instrument")]["trades"] += 1
            instrument_stats[t.get("instrument")]["pnl"] += r["portfolio_pnl"]
        for i, ev in enumerate(r["timeline"], start=1):
            event_rows.append([
                t.get("id"), i, ev.get("type"), ev.get("timestamp"), ev.get("qtyPerAccount"), ev.get("price"),
                ev.get("qtyAfter"), ev.get("avgEntryAfter"), ev.get("eventGross"), ev.get("eventCosts"), ev.get("eventRealized"), ev.get("realizedAfter"), ev.get("reason")
            ])
        for a in t.get("allocations", []):
            allocation_rows.append([
                t.get("id"), a.get("companyName"), a.get("sizeK"), a.get("stage"),
                len(a.get("accountIds", [])), ", ".join(a.get("accountIds", [])), a.get("quantityMultiplier", 1)
            ])

    ws = wb.create_sheet("Trades")
    add_rows(ws,["Trade ID","Date","Instrument","Direction","Status","Target / Account","Buffer pts","Initial Stop Price","Accounts","Gross P&L / Account","Fees & Comm. / Account","Net P&L / Account","Gross Portfolio","Fees & Comm. Portfolio","Net Portfolio P&L","Created UTC","Closed UTC","Close Reason","Manual Reason","Review Adherence","Main Error","Lesson","Reviewed UTC","Notes"],trade_rows)

    ws = wb.create_sheet("Trade Events")
    add_rows(ws,["Trade ID","Seq","Type","Timestamp UTC","Qty / Account","Price","Position After","Average Entry After","Event Gross / Account","Event Fees & Comm. / Account","Event Net / Account","Net Realized After","Reason"],event_rows)

    ws = wb.create_sheet("Allocations")
    add_rows(ws,["Trade ID","Company","Size K","Stage","Account Count","Account IDs","Qty Multiplier"],allocation_rows)

    account_rows = []
    account_balances = {}
    for a in state.get("accounts", []):
        opening = account_opening_balance(a)
        trade_pnl = account_trade_pnl(a.get("id"), trades, replay_cache)
        adjustments = account_adjustments(a)
        current = opening + trade_pnl + adjustments
        account_balances[a.get("id")] = {"opening": opening, "trade_pnl": trade_pnl, "adjustments": adjustments, "current": current}
        account_rows.append([
            a.get("id"), company_map.get(a.get("companyId"), {}).get("name", a.get("companyId")),
            a.get("sizeK"), a.get("stage"), a.get("sequence"), a.get("status"),
            opening, trade_pnl, adjustments, current,
            a.get("createdAt"), a.get("closedAt"), a.get("closureReason"), a.get("replacedById")
        ])
    ws = wb.create_sheet("Accounts")
    add_rows(ws,["Account ID","Company","Size K","Stage","Sequence","Status","Opening Balance","Trading Net P&L","Balance Adjustments","Current Balance","Created UTC","Closed UTC","Closure Reason","Replacement Account ID"],account_rows)

    ws = wb.create_sheet("Companies")
    add_rows(ws,["Company ID","Company Name","Active"],[[c.get("id"), c.get("name"), c.get("active", True)] for c in state.get("companies", [])])

    ws = wb.create_sheet("Daily")
    add_rows(ws,["Date","Trades","Closed Trades","Portfolio P&L"],[[k,v["trades"],v["closed"],v["portfolio_pnl"]] for k,v in sorted(daily.items())])

    ws = wb.create_sheet("Dashboard", 0)
    ws.sheet_view.rightToLeft = True
    ws["A1"] = "Trading OS — Dashboard"
    ws["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="0F172A")
    ws.merge_cells("A1:F1")
    closed = [t for t in trades if t.get("status") == "Closed"]
    total_pnl = sum(replay_cache[t.get("id")]["portfolio_pnl"] for t in closed)
    active_accounts_list = [a for a in state.get("accounts", []) if a.get("status") == "Active"]
    active_accounts = len(active_accounts_list)
    active_balance = sum(account_balances.get(a.get("id"), {}).get("current", 0) for a in active_accounts_list)
    all_balance = sum(x.get("current", 0) for x in account_balances.values())
    active_trades = sum(1 for t in trades if t.get("status") in ("Active", "Partially Closed"))
    kpis = [
        ("Active Balance", active_balance),
        ("All Accounts Balance", all_balance),
        ("Portfolio P&L", total_pnl),
        ("Active Accounts", active_accounts),
        ("Active Trades", active_trades),
        ("Closed Trades", len(closed)),
    ]
    for idx, (label, val) in enumerate(kpis, start=1):
        col = get_column_letter(idx)
        ws[f"{col}3"] = label
        ws[f"{col}3"].font = Font(bold=True, color="94A3B8")
        ws[f"{col}4"] = val
        ws[f"{col}4"].font = Font(size=16, bold=True)
        if label in ("Portfolio P&L", "Active Balance", "All Accounts Balance"):
            ws[f"{col}4"].number_format = '$#,##0.00;[Red]-$#,##0.00'
    ws["A7"] = "Instrument"
    ws["B7"] = "Closed Trades"
    ws["C7"] = "Portfolio P&L"
    for cell in ws[7]:
        if cell.column <= 3:
            cell.fill = PatternFill("solid", fgColor="111827")
            cell.font = Font(color="FFFFFF", bold=True)
    r = 8
    for inst, st in sorted(instrument_stats.items()):
        ws.cell(r,1,inst); ws.cell(r,2,st["trades"]); ws.cell(r,3,st["pnl"]); ws.cell(r,3).number_format='$#,##0.00;[Red]-$#,##0.00'; r += 1
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 20
    if r > 8:
        chart = BarChart(); chart.type = "col"; chart.style = 10; chart.title = "Portfolio P&L by Instrument"; chart.y_axis.title = "P&L"; chart.x_axis.title = "Instrument"
        data = Reference(ws, min_col=3, min_row=7, max_row=r-1)
        cats = Reference(ws, min_col=1, min_row=8, max_row=r-1)
        chart.add_data(data, titles_from_data=True); chart.set_categories(cats); chart.height=7; chart.width=12
        ws.add_chart(chart,"E7")

    for name in ["Trades","Trade Events","Allocations","Accounts","Companies","Daily"]:
        s = wb[name]
        if s.max_row >= 1:
            s.auto_filter.ref = s.dimensions
            for row in s.iter_rows():
                for cell in row:
                    cell.alignment = Alignment(vertical="center", wrap_text=True)

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_PATH)
    print(f"Wrote {OUT_PATH}")


if __name__ == "__main__":
    build()
